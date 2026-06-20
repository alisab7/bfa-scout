"""
Player Registry Excel import — parsing + classification (Player import,
Page 1). Pure functions; the route in `registry_import.py` owns the
upload→preview→commit lifecycle and the parked session.

The registry template (confirmed from real sample files):
  * Sheet  '➡ Player Registry'  (other sheets ignored)
  * Rows 1-2 title banners, row 3 headers, DATA STARTS AT ROW 4
  * A Full Name Arabic | B Full Name (English) | C CPR | D DOB
    E Age Group | F Notes | G # (ignored)
  * Optional (residency files): Nationality + Nationality Code columns,
    matched by HEADER TEXT so citizen files without them still parse.
  * Row 4 is often an EXAMPLE row (Notes contains "Example only").

Rules:
  * CPR → normalize to 9-digit TEXT (born-2000 double-zero handled);
    malformed CPR is a hard ERROR (master key — never imported).
  * Age Group drives FOUR fields via AGE_GROUP_MAP: age_group +
    nationality_status + nationality + nationality_code. Citizens
    (U17/U20/U23/NT/senior) auto-set Bahrain/BHR; residents
    (residency/resident) take nationality + 3-letter code FROM THE FILE
    (required + validated). Unmapped/blank labels → ERROR (no silent default).
  * DOB → mixed datetime/strings; unparseable → import with NULL dob + warn.
  * One player per CPR; highest age group wins (U23>U20>U17>senior) when a
    CPR repeats in-file or already exists in the DB.
  * Idempotent: re-importing the same file produces only UPDATEs, no new rows.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, date

from openpyxl import load_workbook

from app.db import get_db
from app.players.cpr import normalize_cpr


REGISTRY_SHEET = '➡ Player Registry'
DATA_START_ROW = 4                      # rows 1-2 banners, row 3 headers, data row 4+
HEADER_ROW = DATA_START_ROW - 1        # row 3 — column headers
EXAMPLE_MARKER = 'example only'         # Notes substring → skip row

# Highest age group wins on conflict. A player's senior-most squad is their
# ceiling, so a higher rank overrides a lower one.
AGE_GROUP_RANK = {'senior': 0, 'U17': 1, 'U20': 2, 'U23': 3}

# The Age Group cell drives FOUR fields. label → (age_group,
# nationality_status, default_nationality, default_code). Citizen labels
# (U17/U20/U23/NT/senior) auto-set Bahrain/BHR; resident labels leave
# nationality + code None — they MUST come from the file's Nationality /
# Nationality Code columns (required + validated).
#   - 'nt' is the senior National Team (real files use it) → senior+bahraini.
#   - DB never receives 'NT' as age_group — it maps to 'senior' here first.
AGE_GROUP_MAP = {
    'u17':       ('U17',    'bahraini',          'Bahrain', 'BHR'),
    'u20':       ('U20',    'bahraini',          'Bahrain', 'BHR'),
    'u23':       ('U23',    'bahraini',          'Bahrain', 'BHR'),
    'nt':        ('senior', 'bahraini',          'Bahrain', 'BHR'),
    'senior':    ('senior', 'bahraini',          'Bahrain', 'BHR'),
    'residency': ('senior', 'foreign_residency', None,      None),
    'resident':  ('senior', 'foreign_residency', None,      None),
}

NATIONALITY_STATUS_LABELS = {
    'bahraini':          'Citizen (مواطن)',
    'foreign_residency': 'Resident (إقامة)',
}


def map_age_group(raw):
    """Map a raw Age Group cell → (age_group, nationality_status,
    default_nationality, default_code), or None for blank/unmapped labels
    (caller turns None into a clear ERROR — no silent defaults)."""
    if raw is None:
        return None
    return AGE_GROUP_MAP.get(str(raw).strip().lower())


# ───────── DOB parsing ───────────────────────────────────────────────────────

def parse_dob(raw):
    """Return a date, or None if blank/unparseable.

    Handles real Excel datetime cells AND messy strings (leading spaces,
    single-digit day/month). DD/MM/YYYY is the documented format.
    """
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    if not s:
        return None
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None                          # unparseable — caller warns + NULLs


def _clean_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


# ───────── Parse the workbook ────────────────────────────────────────────────

def parse_registry(file_storage) -> list[dict]:
    """Read the '➡ Player Registry' sheet (row 4+) into raw row dicts.

    Uses openpyxl directly (NOT pandas) so cell types survive — the CPR
    rule needs to tell an int (`503061`) from text (`'000503061'`), and the
    DOB rule needs real datetime cells. Raises ValueError on a bad file or
    a missing registry sheet.
    """
    fname = (file_storage.filename or '').lower()
    if not fname.endswith('.xlsx'):
        raise ValueError(f"Unsupported file type {fname!r}. Upload the .xlsx registry.")
    raw = file_storage.read()
    if not raw:
        raise ValueError("Uploaded file is empty.")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read Excel file: {exc}") from exc

    if REGISTRY_SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Sheet {REGISTRY_SHEET!r} not found. Found: {wb.sheetnames}. "
            "Upload the BFA Player Registry template."
        )
    ws = wb[REGISTRY_SHEET]

    # Columns A-F are fixed-position (stable template layout). The two
    # nationality columns are OPTIONAL and matched by HEADER TEXT (row 3),
    # so citizen files without them still parse. "Nationality Code" is
    # matched before plain "Nationality" so the substring doesn't collide.
    nat_col, code_col = None, None
    for col in range(1, ws.max_column + 1):
        h = ws.cell(row=HEADER_ROW, column=col).value
        if h is None:
            continue
        hs = str(h).strip().lower()
        if 'nationality' in hs and 'code' in hs:
            code_col = col
        elif hs == 'nationality':
            nat_col = col

    rows: list[dict] = []
    for excel_row in range(DATA_START_ROW, ws.max_row + 1):
        a = ws.cell(row=excel_row, column=1).value   # Full Name Arabic
        b = ws.cell(row=excel_row, column=2).value   # Full Name English
        c = ws.cell(row=excel_row, column=3).value   # CPR
        d = ws.cell(row=excel_row, column=4).value   # DOB
        e = ws.cell(row=excel_row, column=5).value   # Age Group
        f = ws.cell(row=excel_row, column=6).value   # Notes
        nat  = ws.cell(row=excel_row, column=nat_col).value  if nat_col  else None
        code = ws.cell(row=excel_row, column=code_col).value if code_col else None

        # Skip fully-empty rows (trailing blanks in the sheet).
        if all(v is None or str(v).strip() == ''
               for v in (a, b, c, d, e, f, nat, code)):
            continue

        rows.append({
            'excel_row':           excel_row,
            'name_ar':             _clean_str(a),
            'name_en':             _clean_str(b),
            'cpr_raw':             c,
            'dob_raw':             d,
            'age_group_raw':       e,
            'notes':               _clean_str(f),
            'nationality_raw':     _clean_str(nat),
            'nationality_code_raw': _clean_str(code),
        })
    wb.close()
    return rows


# ───────── Classify ──────────────────────────────────────────────────────────

def _fetch_existing_by_cpr(cprs: list[str]) -> dict:
    """Return {cpr: {'id':, 'age_group':}} for CPRs already in the DB."""
    if not cprs:
        return {}
    conn = get_db()
    with conn.cursor() as cur:
        cur.execute(
            "SELECT id, national_id, age_group FROM players WHERE national_id = ANY(%s)",
            (cprs,)
        )
        return {r['national_id']: {'id': r['id'], 'age_group': r['age_group']}
                for r in cur.fetchall()}


def classify_registry_rows(rows: list[dict]) -> tuple[list[dict], list[dict], dict]:
    """
    Returns (display_rows, write_plan, summary).

    display_rows — one entry PER SOURCE ROW for the preview table:
      {excel_row, name_en, cpr, dob, age_group, status, reason}
      status ∈ NEW | UPDATE | SKIP | ERROR | MERGED
        MERGED = a repeat of a CPR already seen earlier in this file
                 (its data is folded into that CPR's single write).

    write_plan — one entry PER UNIQUE CPR that will actually be written:
      {cpr, full_name, full_name_ar, dob, age_group, db_id}
      (db_id None → INSERT, else UPDATE). age_group already resolved to the
      highest rank across all in-file occurrences AND the existing DB row.

    summary — counts for the preview header.
    """
    # First pass: per-row validation, collect good rows' CPRs.
    prelim: list[dict] = []
    good_cprs: list[str] = []
    for r in rows:
        excel_row = r['excel_row']
        entry = {
            'excel_row': excel_row,
            'name_en':   r['name_en'] or '',
            'cpr':       None,
            'dob':       None,
            'age_group': None,
            'nationality_status': None,
            'nationality':        None,
            'nationality_code':   None,
            'status':    None,
            'reason':    '',
        }

        # Example row → SKIP (checked first so an example with a bad CPR
        # is reported as a skip, not an error).
        if r['notes'] and EXAMPLE_MARKER in r['notes'].lower():
            entry['status'] = 'SKIP'
            entry['reason'] = 'example row'
            prelim.append(entry)
            continue

        cpr = normalize_cpr(r['cpr_raw'])
        if cpr is None:
            entry['status'] = 'ERROR'
            entry['reason'] = f"bad CPR: {r['cpr_raw']!r}"
            prelim.append(entry)
            continue
        entry['cpr'] = cpr

        if not r['name_en']:
            entry['status'] = 'ERROR'
            entry['reason'] = 'missing English name'
            prelim.append(entry)
            continue

        mapped = map_age_group(r['age_group_raw'])
        if mapped is None:
            entry['status'] = 'ERROR'
            entry['reason'] = (
                f"invalid/blank age group: {r['age_group_raw']!r} "
                f"(use U17/U20/U23/NT/senior/residency)")
            prelim.append(entry)
            continue
        age_group, nat_status, def_nat, def_code = mapped
        entry['age_group'] = age_group
        entry['nationality_status'] = nat_status

        # Nationality: citizens auto-get Bahrain/BHR; residents MUST supply
        # nationality + a 3-letter code in the file (no silent default).
        if nat_status == 'foreign_residency':
            nat  = r.get('nationality_raw')
            code = r.get('nationality_code_raw')
            if not nat or not code:
                entry['status'] = 'ERROR'
                entry['reason'] = ("Resident row requires Nationality and "
                                   "Nationality Code columns")
                prelim.append(entry)
                continue
            code = str(code).strip().upper()
            if not re.fullmatch(r'[A-Z]{3}', code):
                entry['status'] = 'ERROR'
                entry['reason'] = (f"Nationality Code must be exactly 3 letters "
                                   f"(got {r.get('nationality_code_raw')!r})")
                prelim.append(entry)
                continue
            entry['nationality'] = str(nat).strip()
            entry['nationality_code'] = code
        else:
            entry['nationality'] = def_nat       # Bahrain
            entry['nationality_code'] = def_code  # BHR

        # DOB: unparseable but non-blank → warn + NULL (don't lose the player).
        dob = parse_dob(r['dob_raw'])
        if dob is None and r['dob_raw'] is not None and str(r['dob_raw']).strip():
            entry['reason'] = f"unreadable DOB {str(r['dob_raw']).strip()!r} → left blank"
        entry['dob'] = dob

        entry['_name_ar'] = r['name_ar']
        entry['_good'] = True
        good_cprs.append(cpr)
        prelim.append(entry)

    # Look up which good CPRs already exist in the DB.
    existing = _fetch_existing_by_cpr(sorted(set(good_cprs)))

    # Second pass: resolve per-CPR write plan with highest-age-wins, and
    # set per-row NEW/UPDATE/MERGED status.
    plan: dict[str, dict] = {}      # cpr -> write entry
    seen_cpr: set[str] = set()
    for entry in prelim:
        if not entry.get('_good'):
            continue
        cpr = entry['cpr']
        db_row = existing.get(cpr)
        db_id = db_row['id'] if db_row else None

        if cpr not in plan:
            # First occurrence of this CPR in the file → it owns the write.
            base_rank = AGE_GROUP_RANK.get(
                db_row['age_group'] if db_row else 'senior', 0)
            row_rank = AGE_GROUP_RANK.get(entry['age_group'], 0)
            best = entry['age_group'] if row_rank >= base_rank else (
                db_row['age_group'] if db_row else 'senior')
            plan[cpr] = {
                'cpr':                cpr,
                'full_name':          entry['name_en'],
                'full_name_ar':       entry.get('_name_ar'),
                'dob':                entry['dob'],
                'age_group':          best,
                'nationality_status': entry['nationality_status'],
                'nationality':        entry['nationality'],
                'nationality_code':   entry['nationality_code'],
                'db_id':              db_id,
            }
            entry['status'] = 'UPDATE' if db_id else 'NEW'
        else:
            # Repeat CPR within the file → fold in (highest age wins,
            # fill missing dob/name_ar), mark the row MERGED.
            p = plan[cpr]
            if AGE_GROUP_RANK.get(entry['age_group'], 0) > AGE_GROUP_RANK.get(p['age_group'], 0):
                p['age_group'] = entry['age_group']
            if p['dob'] is None and entry['dob'] is not None:
                p['dob'] = entry['dob']
            if not p['full_name_ar'] and entry.get('_name_ar'):
                p['full_name_ar'] = entry['_name_ar']
            entry['status'] = 'MERGED'
            entry['reason'] = (entry['reason'] + '; ' if entry['reason'] else '') + \
                              'duplicate CPR in file → merged (highest age group kept)'
        # Reflect the resolved age_group on the display row for clarity.
        entry['age_group'] = plan[cpr]['age_group']

    # Strip internal keys from display rows.
    display_rows = []
    for e in prelim:
        e.pop('_good', None)
        e.pop('_name_ar', None)
        display_rows.append(e)

    write_plan = list(plan.values())
    summary = {
        'total_rows':  len(display_rows),
        'new':         sum(1 for p in write_plan if p['db_id'] is None),
        'updated':     sum(1 for p in write_plan if p['db_id'] is not None),
        'merged':      sum(1 for e in display_rows if e['status'] == 'MERGED'),
        'skipped':     sum(1 for e in display_rows if e['status'] == 'SKIP'),
        'errors':      sum(1 for e in display_rows if e['status'] == 'ERROR'),
        'to_write':    len(write_plan),
    }
    return display_rows, write_plan, summary
